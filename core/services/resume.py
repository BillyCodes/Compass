"""Resume import.

The uploaded file is read in memory, converted to text, scanned for catalog
skills, and discarded. Only the detected skills are kept (ResumeImport.detected),
and nothing reaches the student's Skillprint until the student reviews and
confirms the suggested levels.

Suggested level (PROPOSED rule, documented so it is not hidden in code):
  * mentioned once                         → 2 Beginner
  * mentioned 2+ times                     → 3 Developing
  * mentioned 2+ times and the resume shows
    work/project experience with the skill → 3 (never higher; a resume
    mention is not proof of proficiency, blueprint §6.2)
The student can raise or lower any level before saving.
"""
import io
import re

from django.conf import settings

from .skill_matcher import find_skills

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}
_EXPERIENCE_WORDS = re.compile(
    r"\b(built|developed|implemented|designed|deployed|maintained|led|created|engineered|automated|configured)\b",
    re.IGNORECASE,
)


class ResumeError(Exception):
    """A problem the student can fix (wrong type, too big, unreadable)."""


def _ext(name):
    name = (name or "").lower()
    dot = name.rfind(".")
    return name[dot:] if dot != -1 else ""


def extract_text(uploaded):
    ext = _ext(uploaded.name)
    if ext not in ALLOWED_EXTENSIONS:
        raise ResumeError("Upload a PDF, Word (.docx), or plain-text (.txt) resume.")
    if uploaded.size > settings.COMPASS_MAX_RESUME_BYTES:
        raise ResumeError("That file is larger than 2 MB. Export a smaller copy and try again.")
    data = uploaded.read(settings.COMPASS_MAX_RESUME_BYTES + 1)

    try:
        if ext == ".pdf":
            if not data.startswith(b"%PDF"):
                raise ResumeError("That file doesn't look like a PDF.")
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise ResumeError("That PDF is password-protected. Upload an unlocked copy.")
            text = "\n".join((page.extract_text() or "") for page in reader.pages[:10])
        elif ext == ".docx":
            if not data.startswith(b"PK"):
                raise ResumeError("That file doesn't look like a Word document.")
            import docx

            document = docx.Document(io.BytesIO(data))
            parts = [p.text for p in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    parts.extend(cell.text for cell in row.cells)
            text = "\n".join(parts)
        else:
            text = data.decode("utf-8", errors="replace")
    except ResumeError:
        raise
    except Exception:
        raise ResumeError("We couldn't read that file. Try saving it again as PDF or .docx.")

    text = text[: settings.COMPASS_MAX_TEXT_CHARS * 2]
    if len(text.strip()) < 40:
        raise ResumeError(
            "We couldn't find text in that file. If it's a scanned image, upload a text-based PDF or .docx."
        )
    return text


def suggest_level(hit, text):
    level = 2 if hit.hits == 1 else 3
    return level


def detect_skills(text):
    """Return a list of suggestion dicts sorted by strength of evidence."""
    hits = find_skills(text)
    suggestions = []
    for skill_id, hit in hits.items():
        experienced = any(_EXPERIENCE_WORDS.search(c) for c in hit.contexts)
        suggestions.append(
            {
                "skill_id": skill_id,
                "name": hit.skill.name,
                "category": hit.skill.category,
                "hits": hit.hits,
                "suggested_level": suggest_level(hit, text),
                "experience": bool(experienced),
                "context": hit.contexts[0] if hit.contexts else "",
            }
        )
    suggestions.sort(key=lambda s: (-s["hits"], s["name"]))
    return suggestions
